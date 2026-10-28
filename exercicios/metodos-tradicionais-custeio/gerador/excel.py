"""Planilha de resolução (layout do modelo 'Planilha Resolução Metodos Tradicionais.xlsx'), uma aba por exercício, com fórmulas."""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1]
EXPECT = []  # (aba, célula, valor esperado) para conferência após recálculo

RS = '_-"R$"* #,##0.00_-;\\-"R$"* #,##0.00_-;_-"R$"* "-"??_-;_-@'
PCT1 = "0.0%"
PCT2 = "0.00%"
PP = "0.00"      # pontos percentuais (ex.: 27,60)
DEC = "0.00"
INT = "#,##0"

NAVY, GREEN, BLUE, DKGREEN, LBLUE, CYAN = "FF1F3864", "FF00B050", "FF4472C4", "FF385623", "FFB4C6E7", "FF00B0F0"
thin, med, dbl = Side(style="thin"), Side(style="medium"), Side(style="double")


def fill(c):
    return PatternFill("solid", fgColor=c)


def box(ws, rng, side=thin):
    for row in ws[rng]:
        for c in row:
            c.border = Border(left=side, right=side, top=side, bottom=side)


def put(ws, ref, v, fmt=None, bold=False, color=None, bg=None, size=11, al=None, font="Arial", wrap=False, italic=False):
    c = ws[ref]
    c.value = v
    c.font = Font(name=font, size=size, bold=bold, color=color, italic=italic)
    if fmt:
        c.number_format = fmt
    if bg:
        c.fill = fill(bg)
    if al or wrap:
        c.alignment = Alignment(horizontal=al, vertical="center", wrap_text=wrap)
    return c


def exp(ws, ref, v):
    EXPECT.append((ws.title, ref, float(v)))


def title(ws, rng, text, bg=GREEN, color="FFFFFFFF", size=13):
    ws.merge_cells(rng)
    a = rng.split(":")[0]
    put(ws, a, text, bold=True, color=color, bg=bg, size=size, al="center", font="Calibri")
    box(ws, rng, med)


def header(ws, row, col0, labels):
    for i, t in enumerate(labels):
        put(ws, f"{L(col0 + i)}{row}", t, bold=True, bg=LBLUE, size=11, al="center", font="Calibri", wrap=True)
        ws[f"{L(col0 + i)}{row}"].border = Border(left=thin, right=thin, top=thin, bottom=dbl)
    ws.row_dimensions[row].height = 30


# ---------------- DRE (colunas B:D, como no modelo) ----------------
def dre(ws, empresa, R, cvv, cvp, cf):
    """cvv/cvp/cf: listas (rótulo, valor); cvp[0] é a matéria-prima. Retorna dicionário de linhas."""
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 52
    ws.column_dimensions["C"].width = 19
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 3
    ws.merge_cells("B4:D4")
    put(ws, "B4", "DEMONSTRATIVO RESULTADOS DO EXERCICIO", bold=True, color="FFF2F2F2", bg=NAVY, size=14, al="center")
    ws.merge_cells("B5:D5")
    put(ws, "B5", empresa, bold=True, size=11, al="center", italic=True)
    for a, t in (("B6", "DISCRIMINAÇÃO"), ("C6", "VALOR R$"), ("D6", "%")):
        put(ws, a, t, bold=True, color="FFF2F2F2", bg=GREEN if a != "D6" else BLUE, al="center")
    box(ws, "B4:D6")
    r = {}

    def sec(row, label, formula, bg=BLUE):
        put(ws, f"B{row}", label, bold=True, color="FFF2F2F2", bg=bg)
        put(ws, f"C{row}", formula, RS, bold=True, font="Calibri", size=12)
        put(ws, f"D{row}", f"=C{row}/$C$8", PCT1, bold=True, font="Calibri", size=12)
        for col in "BCD":
            ws[f"{col}{row}"].border = Border(left=thin, right=thin, top=med, bottom=med)

    def items(row0, lst):
        for i, (lab, v) in enumerate(lst):
            rr = row0 + i
            put(ws, f"B{rr}", lab)
            put(ws, f"C{rr}", v, RS, font="Calibri", size=11)
            put(ws, f"D{rr}", f"=C{rr}/$C$8", PCT2, font="Calibri", size=11)
            box(ws, f"B{rr}:D{rr}")
        return row0 + len(lst) - 1

    sec(8, "1. Receita Total", R)
    ws["D8"].value = "=C8/$C$8"
    r["R"] = 8
    row = 10
    last = items(row + 1, cvv)
    sec(row, "2. Custos Variáveis de Vendas", f"=SUM(C{row + 1}:C{last})")
    r["CVV"] = row
    row = last + 2
    last = items(row + 1, cvp)
    sec(row, "3. Custos Variáveis de Produção", f"=SUM(C{row + 1}:C{last})")
    r["CVP"], r["MP"] = row, row + 1
    row = last + 2
    sec(row, "4. Custos Variáveis Totais", f"=C{r['CVV']}+C{r['CVP']}")
    r["CVT"] = row
    row += 2
    sec(row, "5. Margem de Contribuição", f"=C8-C{r['CVT']}")
    r["MC"] = row
    row += 2
    last = items(row + 1, cf)
    sec(row, "6. Custos Fixos Totais", f"=SUM(C{row + 1}:C{last})")
    r["CF"] = row
    r["cf_items"] = {lab: row + 1 + i for i, (lab, _) in enumerate(cf)}
    row = last + 2
    sec(row, "7. Resultado Operacional", f"=C{r['MC']}-C{r['CF']}", bg=DKGREEN)
    r["LO"] = row
    # percentuais para as fórmulas
    row += 3
    title(ws, f"B{row}:C{row}", "PERCENTUAIS PARA AS FÓRMULAS (em pontos %)", bg=CYAN, size=12)
    put(ws, f"B{row + 1}", "CVP sem matéria-prima = CVP% − MP%")
    put(ws, f"C{row + 1}", f"=(D{r['CVP']}-D{r['MP']})*100", PP, font="Calibri", size=12)
    put(ws, f"B{row + 2}", "CVV")
    put(ws, f"C{row + 2}", f"=D{r['CVV']}*100", PP, font="Calibri", size=12)
    put(ws, f"B{row + 3}", "CF (usado só no custeio marginal)")
    put(ws, f"C{row + 3}", f"=D{r['CF']}*100", PP, font="Calibri", size=12)
    put(ws, f"B{row + 4}", "Margem de contribuição")
    put(ws, f"C{row + 4}", f"=D{r['MC']}*100", PP, font="Calibri", size=12)
    box(ws, f"B{row + 1}:C{row + 4}")
    r["pCVP"], r["pCVV"], r["pCF"], r["pMC"] = (f"$C${row + 1}", f"$C${row + 2}", f"$C${row + 3}", f"$C${row + 4}")
    r["end"] = row + 4
    return r


def setcols(ws):
    for col, w in zip("FGHIJKLMNO", (3, 3, 34, 14, 18, 16, 16, 16, 10, 17)):
        ws.column_dimensions[col].width = w


# ---------------- tabelas de métodos (colunas H:O, como na aba "Questão 3") ----------------
def produtos(ws, row, prods, extra=None):
    """prods: (nome, mp, produzidas, vendidas, lucro[, mercado]). Retorna dict de referências."""
    cols = ["PRODUTO", "MP (R$/un.)", "PRODU-ZIDAS", "VENDI-DAS", "LUCRO (%)"] + (["PREÇO DE MERCADO"] if extra else [])
    title(ws, f"H{row}:{L(7 + len(cols))}{row}", "DADOS DOS PRODUTOS", bg=CYAN)
    header(ws, row + 1, 8, cols)
    ref = {"rows": []}
    for i, p in enumerate(prods):
        rr = row + 2 + i
        put(ws, f"H{rr}", p[0], font="Calibri", size=12)
        put(ws, f"I{rr}", p[1], DEC, font="Calibri", size=12)
        put(ws, f"J{rr}", p[2], INT, font="Calibri", size=12)
        put(ws, f"K{rr}", p[3], INT, font="Calibri", size=12)
        put(ws, f"L{rr}", p[4], "0", font="Calibri", size=12)
        if extra:
            put(ws, f"M{rr}", p[5], RS, font="Calibri", size=12)
        ref["rows"].append(rr)
    tr = row + 2 + len(prods)
    put(ws, f"H{tr}", "TOTAL", bold=True, font="Calibri", size=12)
    put(ws, f"J{tr}", f"=SUM(J{row + 2}:J{tr - 1})", INT, bold=True, font="Calibri", size=12)
    put(ws, f"K{tr}", f"=SUM(K{row + 2}:K{tr - 1})", INT, bold=True, font="Calibri", size=12)
    box(ws, f"H{row + 2}:{L(7 + len(cols))}{tr}")
    ref["tot"] = tr
    return ref


def metodo(ws, row, kind, d, items, cfu=None):
    """kind: 'mar' | 'int' | 'ide'. items: lista de (célula nome, célula MP, célula lucro). Retorna linhas."""
    nome = {"mar": "CUSTEIO MARGINAL", "int": "CUSTEIO POR ABSORÇÃO INTEGRAL", "ide": "CUSTEIO POR ABSORÇÃO IDEAL"}[kind]
    title(ws, f"H{row}:O{row}", nome)
    header(ws, row + 1, 8, ["PRODUTO", "MP", "CVV", "CVP", "CF (%)" if kind == "mar" else "CF UNITARIO", "CUSTO TOTAL", "LUCRO", "P. VENDA"])
    out = []
    for i, (cn, cmp, cl) in enumerate(items):
        rr = row + 2 + i
        put(ws, f"H{rr}", f"={cn}", font="Calibri", size=12)
        put(ws, f"I{rr}", f"={cmp}", DEC, font="Calibri", size=12)
        put(ws, f"J{rr}", f"={d['pCVV']}", PP, font="Calibri", size=12)
        put(ws, f"K{rr}", f"={d['pCVP']}", PP, font="Calibri", size=12)
        if kind == "mar":
            put(ws, f"L{rr}", f"={d['pCF']}", PP, font="Calibri", size=12)
            put(ws, f"M{rr}", f"=I{rr}/((100-(J{rr}+K{rr}+L{rr}))/100)", RS, font="Calibri", size=12)
            put(ws, f"O{rr}", f"=I{rr}/((100-(J{rr}+K{rr}+L{rr}+N{rr}))/100)", RS, bold=True, font="Calibri", size=12)
        else:
            put(ws, f"L{rr}", f"={cfu}", RS, font="Calibri", size=12)
            put(ws, f"M{rr}", f"=(I{rr}+L{rr})/((100-(J{rr}+K{rr}))/100)", RS, font="Calibri", size=12)
            put(ws, f"O{rr}", f"=(I{rr}+L{rr})/((100-(J{rr}+K{rr}+N{rr}))/100)", RS, bold=True, font="Calibri", size=12)
        put(ws, f"N{rr}", f"={cl}", "0", font="Calibri", size=12)
        out.append(rr)
    box(ws, f"H{row + 2}:O{row + 1 + len(items)}")
    note = ("CUSTO = MP ÷ [(100 − (CVV+CVP+CF))/100]   ·   PREÇO = MP ÷ [(100 − (CVV+CVP+CF+L))/100]" if kind == "mar" else
            "CUSTO = (MP + CF unit.) ÷ [(100 − (CVV+CVP))/100]   ·   PREÇO = (MP + CF unit.) ÷ [(100 − (CVV+CVP+L))/100]")
    put(ws, f"H{row + 2 + len(items)}", note, italic=True, size=9, color="FF595959")
    return out


def pe_block(ws, row, d, dep_ref, ativo, taxa):
    title(ws, f"H{row}:J{row}", "DADOS · PONTOS DE EQUILÍBRIO", bg=CYAN)
    lines = [("RECEITA", f"=C{d['R']}", RS), ("MARGEM DE CONTRIBUIÇÃO", f"=C{d['MC']}", RS), ("MC %", f"=D{d['MC']}", PCT2),
             ("CUSTO FIXO", f"=C{d['CF']}", RS), ("DEPRECIAÇÃO", f"={dep_ref}", RS), ("ATIVO TOTAL", ativo, RS),
             ("TAXA DE OPORTUNIDADE (a.m.)", taxa, PCT2)]
    for i, (lab, v, f) in enumerate(lines):
        rr = row + 1 + i
        ws.merge_cells(f"H{rr}:I{rr}")
        put(ws, f"H{rr}", lab, bold=True, font="Calibri", size=12)
        put(ws, f"J{rr}", v, f, font="Calibri", size=12)
    co = row + 8
    ws.merge_cells(f"H{co}:I{co}")
    put(ws, f"H{co}", "CUSTO DE OPORTUNIDADE = ATIVO × TAXA", bold=True, font="Calibri", size=12)
    put(ws, f"J{co}", f"=J{row + 6}*J{row + 7}", RS, font="Calibri", size=12)
    box(ws, f"H{row + 1}:J{co}")
    R_, MC_, CF_, DEP_, CO_ = f"J{row + 1}", f"J{row + 3}", f"J{row + 4}", f"J{row + 5}", f"J{co}"
    res = [("PONTO DE EQUILÍBRIO FINANCEIRO = (CF − Depreciação) ÷ MC%", f"=({CF_}-{DEP_})/{MC_}", RS, "pef"),
           ("PONTO DE EQUILÍBRIO CONTÁBIL = CF ÷ MC%", f"={CF_}/{MC_}", RS, "pec"),
           ("PONTO DE EQUILÍBRIO ECONÔMICO = (CF + Custo de oportunidade) ÷ MC%", f"=({CF_}+{CO_})/{MC_}", RS, "pee")]
    refs = {}
    for i, (lab, f, fm, k) in enumerate(res):
        rr = co + 2 + i
        ws.merge_cells(f"H{rr}:I{rr}")
        put(ws, f"H{rr}", lab, bold=True, font="Calibri", size=11, wrap=True)
        put(ws, f"J{rr}", f, fm, bold=True, font="Calibri", size=12)
        ws.row_dimensions[rr].height = 32
        refs[k] = f"J{rr}"
    box(ws, f"H{co + 2}:J{co + 4}", med)
    refs["R"], refs["next"] = R_, co + 6
    return refs


def resposta(ws, row, pergunta, texto, height=60):
    ws.merge_cells(f"H{row}:O{row}")
    put(ws, f"H{row}", pergunta, bold=True, font="Calibri", size=12, color="FFFF0000")
    ws.merge_cells(f"H{row + 1}:O{row + 1}")
    put(ws, f"H{row + 1}", texto, font="Calibri", size=11, wrap=True, al="left")
    ws.row_dimensions[row + 1].height = height
    box(ws, f"H{row + 1}:O{row + 1}")
    return row + 3


def enunciado(ws, txt):
    ws.merge_cells("B2:O2")
    put(ws, "B2", txt, bold=True, size=12, color="FF1F3864")


wb = openpyxl.Workbook()
wb.remove(wb.active)

# ===================== EXERCÍCIO 1 =====================
ws = wb.create_sheet("Exercício 1")
enunciado(ws, "EXERCÍCIO 1 · Laticínios Serra Verde · Classificação, DRE, pontos de equilíbrio e preço pelo custeio marginal")
setcols(ws)
d = dre(ws, "Laticínios Serra Verde Ltda. · março",
        480000,
        [("ICMS sobre vendas", 57600), ("PIS/COFINS sobre vendas", 14400), ("Comissão dos vendedores", 14400), ("Frete de entrega aos supermercados", 9600)],
        [("Leite in natura (matéria-prima)", 168000), ("Embalagens primárias (potes, filmes, rótulos)", 19200), ("Energia elétrica da fábrica", 14400),
         ("Prêmio de produtividade da produção", 4800), ("Frete sobre compras de leite", 9600)],
        [("Salários da fábrica (mensalistas) + encargos", 38000), ("Salários da administração + encargos", 22000), ("Pró-labore dos sócios", 17000),
         ("Depreciação de máquinas e câmaras frias", 12000), ("Aluguel do centro de distribuição", 8000), ("Seguros", 2400), ("Telefone e internet", 1600),
         ("Honorários contábeis", 3000), ("Material de expediente", 500), ("Manutenção predial (contrato mensal fixo)", 2500), ("IPTU e taxas municipais", 1000)])
for k, v in (("CVV", 96000), ("CVP", 216000), ("CVT", 312000), ("MC", 168000), ("CF", 108000), ("LO", 60000)):
    exp(ws, f"C{d[k]}", v)
put(ws, "B3", "a) A classificação está na própria estrutura do DRE: itens do grupo 2 = CVV, grupo 3 = CVP, grupo 6 = CF.", italic=True, size=10, color="FF595959")
pe = pe_block(ws, 5, d, f"C{d['cf_items']['Depreciação de máquinas e câmaras frias']}", 1800000, 0.015)
exp(ws, pe["pef"], 274285.71); exp(ws, pe["pec"], 308571.43); exp(ws, pe["pee"], 385714.29)
r = pe["next"]
ws.merge_cells(f"H{r}:I{r}")
put(ws, f"H{r}", "MARGEM DE SEGURANÇA = (Receita − PE contábil) ÷ Receita", bold=True, font="Calibri", size=11, wrap=True)
put(ws, f"J{r}", f"=({pe['R']}-{pe['pec']})/{pe['R']}", PCT2, bold=True, font="Calibri", size=12)
ws.row_dimensions[r].height = 32
box(ws, f"H{r}:J{r}", med)
exp(ws, f"J{r}", 0.357143)
r += 3
title(ws, f"H{r}:L{r}", "e) DADOS · QUEIJO MINAS FRESCAL", bg=CYAN)
header(ws, r + 1, 8, ["PRODUTO", "MP", "", "", "LUCRO (%)"])
put(ws, f"H{r + 2}", "Queijo minas frescal (500 g)", font="Calibri", size=12)
put(ws, f"I{r + 2}", 9.80, DEC, font="Calibri", size=12)
put(ws, f"L{r + 2}", 10, "0", font="Calibri", size=12)
box(ws, f"H{r + 2}:L{r + 2}")
rows = metodo(ws, r + 4, "mar", d, [(f"H{r + 2}", f"I{r + 2}", f"L{r + 2}")])
exp(ws, f"M{rows[0]}", 20.63); exp(ws, f"O{rows[0]}", 26.13)

# ===================== EXERCÍCIO 2 =====================
ws = wb.create_sheet("Exercício 2")
enunciado(ws, "EXERCÍCIO 2 · Metalúrgica Ferraço · Preço de venda pelo custeio marginal")
setcols(ws)
d = dre(ws, "Metalúrgica Ferraço Ltda. · setembro", 350000,
        [("Custos variáveis de vendas (impostos, comissões, frete de entrega)", 63000)],
        [("Matéria-prima (aço, latão, zamac)", 105000), ("Outros CVP (energia, insumos, frete sobre compras)", 35000)],
        [("Custos fixos", 77000)])
exp(ws, f"C{d['LO']}", 70000)
title(ws, "H5:L5", "DADOS DOS ITENS", bg=CYAN)
header(ws, 6, 8, ["ITEM", "MP (R$/un.)", "", "", "LUCRO (%)"])
its = [("Dobradiça 3\" (par)", 8.40, 10), ("Fechadura de embutir", 46.50, 15), ("Puxador de inox 30 cm", 12.90, 8), ("Trinco de sobrepor", 5.60, 20)]
for i, (n, mp, l) in enumerate(its):
    put(ws, f"H{7 + i}", n, font="Calibri", size=12); put(ws, f"I{7 + i}", mp, DEC, font="Calibri", size=12); put(ws, f"L{7 + i}", l, "0", font="Calibri", size=12)
box(ws, "H7:L10")
rows = metodo(ws, 13, "mar", d, [(f"H{7 + i}", f"I{7 + i}", f"L{7 + i}") for i in range(4)])
for rr, (c, p) in zip(rows, [(16.80, 21.00), (93.00, 132.86), (25.80, 30.71), (11.20, 18.67)]):
    exp(ws, f"M{rr}", c); exp(ws, f"O{rr}", p)
r = rows[-1] + 3
title(ws, f"H{r}:K{r}", "c) PREÇO DO CONCORRENTE · FECHADURA", bg=CYAN)
lines = [("Preço do concorrente", 120, RS), ("MP da fechadura", f"=I{rows[1]}", DEC), ("MP ÷ Preço (em %)", f"=J{r + 2}/J{r + 1}*100", PP),
         ("Margem obtida L = 100 − (CVV+CVP+CF) − MP/Preço", f"=100-({d['pCVV']}+{d['pCVP']}+{d['pCF']})-J{r + 3}", PP)]
for i, (lab, v, f) in enumerate(lines):
    ws.merge_cells(f"H{r + 1 + i}:I{r + 1 + i}")
    put(ws, f"H{r + 1 + i}", lab, bold=True, font="Calibri", size=11, wrap=True)
    put(ws, f"J{r + 1 + i}", v, f, font="Calibri", size=12, bold=(i == 3))
box(ws, f"H{r + 1}:J{r + 4}")
ws.row_dimensions[r + 4].height = 32
exp(ws, f"J{r + 4}", 11.25)
r += 6
title(ws, f"H{r}:K{r}", "d) DIVISOR DO PREÇO COM L = 50%", bg=CYAN)
ws.merge_cells(f"H{r + 1}:I{r + 1}")
put(ws, f"H{r + 1}", "(100 − (CVV+CVP+CF+50))/100", bold=True, font="Calibri", size=11)
put(ws, f"J{r + 1}", f"=(100-({d['pCVV']}+{d['pCVP']}+{d['pCF']}+50))/100", DEC, font="Calibri", size=12, bold=True)
box(ws, f"H{r + 1}:J{r + 1}")
exp(ws, f"J{r + 1}", 0)
resposta(ws, r + 3, "RESPOSTA d)", "Com L = 50% o divisor fica zero e não existe divisão por zero: o preço tenderia ao infinito. "
         "CVV + CVP + CF já somam 50% da receita, então a soma com o lucro precisa ficar abaixo de 100%.")

# ===================== EXERCÍCIO 3 =====================
ws = wb.create_sheet("Exercício 3")
enunciado(ws, "EXERCÍCIO 3 · Móveis Pinheiral · Absorção integral × ideal com produção maior que vendas")
setcols(ws)
d = dre(ws, "Móveis Pinheiral Ltda. · outubro", 600000,
        [("Custos variáveis de vendas", 120000)],
        [("Matéria-prima", 210000), ("Outros CVP", 60000)],
        [("Custos fixos", 144000)])
exp(ws, f"C{d['LO']}", 66000)
pr = produtos(ws, 5, [("Cômoda 4 gavetas", 120, 800, 600, 12), ("Mesa de jantar 6 lugares", 210, 500, 400, 15), ("Cadeira estofada", 45, 1700, 1000, 10)])
t = pr["tot"]
r = t + 2
title(ws, f"H{r}:J{r}", "b) CUSTO FIXO UNITÁRIO", bg=CYAN)
ws.merge_cells(f"H{r + 1}:I{r + 1}"); ws.merge_cells(f"H{r + 2}:I{r + 2}")
put(ws, f"H{r + 1}", "INTEGRAL = CF ÷ Σ itens produzidos", bold=True, font="Calibri", size=11)
put(ws, f"J{r + 1}", f"=ROUND(C{d['CF']}/J{t},2)", RS, font="Calibri", size=12, bold=True)
put(ws, f"H{r + 2}", "IDEAL = CF ÷ Σ itens vendidos", bold=True, font="Calibri", size=11)
put(ws, f"J{r + 2}", f"=ROUND(C{d['CF']}/K{t},2)", RS, font="Calibri", size=12, bold=True)
box(ws, f"H{r + 1}:J{r + 2}")
exp(ws, f"J{r + 1}", 48); exp(ws, f"J{r + 2}", 72)
cfi, cfd = f"$J${r + 1}", f"$J${r + 2}"
it = [(f"$H${x}", f"$I${x}", f"$L${x}") for x in pr["rows"]]
ri = metodo(ws, r + 4, "int", d, it, cfi)
rd = metodo(ws, ri[-1] + 3, "ide", d, it, cfd)
for rr, (c, p) in zip(ri, [(240.00, 289.66), (368.57, 469.09), (132.86, 155.00)]):
    exp(ws, f"M{rr}", c); exp(ws, f"O{rr}", p)
for rr, (c, p) in zip(rd, [(274.29, 331.03), (402.86, 512.73), (167.14, 195.00)]):
    exp(ws, f"M{rr}", c); exp(ws, f"O{rr}", p)
r = rd[-1] + 3
title(ws, f"H{r}:K{r}", "COMPARAÇÃO · PREÇO IDEAL × INTEGRAL", bg=CYAN)
header(ws, r + 1, 8, ["PRODUTO", "PREÇO INTEGRAL", "PREÇO IDEAL", "DIFERENÇA"])
for i, (a, b) in enumerate(zip(ri, rd)):
    rr = r + 2 + i
    put(ws, f"H{rr}", f"=H{a}", font="Calibri", size=12)
    put(ws, f"I{rr}", f"=O{a}", RS, font="Calibri", size=12)
    put(ws, f"J{rr}", f"=O{b}", RS, font="Calibri", size=12)
    put(ws, f"K{rr}", f"=J{rr}/I{rr}-1", PCT1, font="Calibri", size=12)
box(ws, f"H{r + 2}:K{r + 4}")
r += 6
title(ws, f"H{r}:J{r}", "d) CUSTO FIXO QUE FICA NO ESTOQUE (INTEGRAL)", bg=CYAN)
ws.merge_cells(f"H{r + 1}:I{r + 1}")
put(ws, f"H{r + 1}", "CF unit. integral × (produzidas − vendidas)", bold=True, font="Calibri", size=11)
put(ws, f"J{r + 1}", f"={cfi}*(J{t}-K{t})", RS, font="Calibri", size=12, bold=True)
box(ws, f"H{r + 1}:J{r + 1}")
exp(ws, f"J{r + 1}", 48000)
resposta(ws, r + 3, "RESPOSTA e)", "Num mês de formação de estoque, o preço pela absorção IDEAL é o mais prudente: as vendas do mês pagam todo o CF. "
         "O integral deixa parte do CF no estoque e só é seguro se o estoque da feira for vendido. Limitação: dividir o CF pela contagem de peças "
         "dá à cadeira o mesmo CF por unidade que a mesa (motivação para UEP e ABC).", 75)

# ===================== EXERCÍCIO 4 =====================
ws = wb.create_sheet("Exercício 4")
enunciado(ws, "EXERCÍCIO 4 · Confecções Vale do Itajaí · Os três métodos lado a lado")
setcols(ws)
d = dre(ws, "Confecções Vale do Itajaí Ltda. · agosto", 320000,
        [("Comissão dos representantes", 16000), ("Frete de entrega", 9600), ("Impostos sobre a venda", 32000)],
        [("Tecido e aviamentos (matéria-prima)", 96000), ("Energia elétrica da fábrica", 12800), ("Linhas, etiquetas e embalagens", 6400)],
        [("Salário da produção (mensalistas)", 28000), ("Salário da administração", 14000), ("Depreciação", 4000), ("Aluguel do galpão", 6000),
         ("Outras despesas fixas", 12000)])
for k, v in (("CVV", 57600), ("CVP", 115200), ("CVT", 172800), ("MC", 147200), ("CF", 64000), ("LO", 83200)):
    exp(ws, f"C{d[k]}", v)
pr = produtos(ws, 5, [("Camiseta básica", 9.60, 4000, 3600, 15), ("Calça jeans", 38.50, 2000, 1700, 20), ("Jaqueta jeans", 72.00, 600, 500, 25),
                      ("Vestido midi", 44.00, 1400, 1000, 18), ("Bermuda sarja", 21.00, 2000, 1200, 12)])
t = pr["tot"]
r = t + 2
title(ws, f"H{r}:J{r}", "CUSTO FIXO UNITÁRIO", bg=CYAN)
ws.merge_cells(f"H{r + 1}:I{r + 1}"); ws.merge_cells(f"H{r + 2}:I{r + 2}")
put(ws, f"H{r + 1}", "INTEGRAL = CF ÷ Σ itens produzidos", bold=True, font="Calibri", size=11)
put(ws, f"J{r + 1}", f"=ROUND(C{d['CF']}/J{t},2)", RS, font="Calibri", size=12, bold=True)
put(ws, f"H{r + 2}", "IDEAL = CF ÷ Σ itens vendidos", bold=True, font="Calibri", size=11)
put(ws, f"J{r + 2}", f"=ROUND(C{d['CF']}/K{t},2)", RS, font="Calibri", size=12, bold=True)
box(ws, f"H{r + 1}:J{r + 2}")
exp(ws, f"J{r + 1}", 6.40); exp(ws, f"J{r + 2}", 8.00)
cfi, cfd = f"$J${r + 1}", f"$J${r + 2}"
it = [(f"$H${x}", f"$I${x}", f"$L${x}") for x in pr["rows"]]
rm = metodo(ws, r + 4, "mar", d, it)
ri = metodo(ws, rm[-1] + 3, "int", d, it, cfi)
rd = metodo(ws, ri[-1] + 3, "ide", d, it, cfd)
gab4 = [((17.14, 23.41), (21.05, 26.23), (23.16, 28.85)), ((68.75, 106.94), (59.08, 80.18), (61.18, 83.04)),
        ((128.57, 232.26), (103.16, 153.73), (105.26, 156.86)), ((78.57, 115.79), (66.32, 86.90), (68.42, 89.66)),
        ((37.50, 47.73), (36.05, 42.81), (38.16, 45.31))]
for i, g in enumerate(gab4):
    for rows_, (c, p) in zip((rm, ri, rd), g):
        exp(ws, f"M{rows_[i]}", c); exp(ws, f"O{rows_[i]}", p)
r = rd[-1] + 3
title(ws, f"H{r}:L{r}", "c) COMPARAÇÃO DOS PREÇOS", bg=CYAN)
header(ws, r + 1, 8, ["PRODUTO", "MARGINAL", "INTEGRAL", "IDEAL", "MAIOR ÷ MENOR"])
for i in range(5):
    rr = r + 2 + i
    put(ws, f"H{rr}", f"=H{rm[i]}", font="Calibri", size=12)
    put(ws, f"I{rr}", f"=O{rm[i]}", RS, font="Calibri", size=12)
    put(ws, f"J{rr}", f"=O{ri[i]}", RS, font="Calibri", size=12)
    put(ws, f"K{rr}", f"=O{rd[i]}", RS, font="Calibri", size=12)
    put(ws, f"L{rr}", f"=MAX(I{rr}:K{rr})/MIN(I{rr}:K{rr})-1", PCT1, font="Calibri", size=12)
box(ws, f"H{r + 2}:L{r + 6}")
r += 8
title(ws, f"H{r}:J{r}", "MP EM QUE CUSTO MARGINAL = CUSTO INTEGRAL", bg=CYAN)
lines = [("Divisor marginal (100 − (CVV+CVP+CF))/100", f"=(100-({d['pCVV']}+{d['pCVP']}+{d['pCF']}))/100", DEC),
         ("Divisor absorção (100 − (CVV+CVP))/100", f"=(100-({d['pCVV']}+{d['pCVP']}))/100", DEC),
         ("MP de virada = CFu × Div.marg ÷ (Div.abs − Div.marg)", f"={cfi}*J{r + 1}/(J{r + 2}-J{r + 1})", RS)]
for i, (lab, v, f) in enumerate(lines):
    ws.merge_cells(f"H{r + 1 + i}:I{r + 1 + i}")
    put(ws, f"H{r + 1 + i}", lab, bold=True, font="Calibri", size=11, wrap=True)
    put(ws, f"J{r + 1 + i}", v, f, font="Calibri", size=12, bold=(i == 2))
    ws.row_dimensions[r + 1 + i].height = 30
box(ws, f"H{r + 1}:J{r + 3}")
exp(ws, f"J{r + 3}", 17.92)
resposta(ws, r + 5, "RESPOSTA c)", "O marginal dá o maior preço para calça, jaqueta, vestido e bermuda (MP acima de R$ 17,92) e o menor para a camiseta. "
         "No marginal o CF é distribuído em proporção à MP; na absorção, todas as peças recebem o mesmo CF em reais.", 50)

# ===================== EXERCÍCIO 5 =====================
ws = wb.create_sheet("Exercício 5")
enunciado(ws, "EXERCÍCIO 5 · Polimix Embalagens · Preço, mercado e a espiral da morte")
setcols(ws)
d = dre(ws, "Polimix Embalagens Plásticas Ltda. · julho", 250000,
        [("Custos variáveis de vendas", 37500)],
        [("Resina PP (matéria-prima)", 100000), ("Outros CVP", 25000)],
        [("Depreciação das injetoras", 10000), ("Outros custos fixos", 40000)])
exp(ws, f"C{d['LO']}", 37500)
pe = pe_block(ws, 5, d, f"C{d['cf_items']['Depreciação das injetoras']}", 1200000, 0.0125)
exp(ws, pe["pef"], 114285.71); exp(ws, pe["pec"], 142857.14); exp(ws, pe["pee"], 185714.29)
r = pe["next"]
pr = produtos(ws, r, [("Balde 20 L", 6.30, 6500, 5000, 10, 19.90), ("Bombona 50 L", 18.40, 2000, 1500, 10, 49.90),
                      ("Caixa organizadora 30 L", 9.10, 4000, 3500, 10, 24.90)], extra=True)
t = pr["tot"]
r = t + 2
title(ws, f"H{r}:J{r}", "CUSTO FIXO UNITÁRIO", bg=CYAN)
ws.merge_cells(f"H{r + 1}:I{r + 1}"); ws.merge_cells(f"H{r + 2}:I{r + 2}")
put(ws, f"H{r + 1}", "INTEGRAL = CF ÷ Σ itens produzidos", bold=True, font="Calibri", size=11)
put(ws, f"J{r + 1}", f"=ROUND(C{d['CF']}/J{t},2)", RS, font="Calibri", size=12, bold=True)
put(ws, f"H{r + 2}", "IDEAL = CF ÷ Σ itens vendidos", bold=True, font="Calibri", size=11)
put(ws, f"J{r + 2}", f"=ROUND(C{d['CF']}/K{t},2)", RS, font="Calibri", size=12, bold=True)
box(ws, f"H{r + 1}:J{r + 2}")
exp(ws, f"J{r + 1}", 4); exp(ws, f"J{r + 2}", 5)
cfi, cfd = f"$J${r + 1}", f"$J${r + 2}"
it = [(f"$H${x}", f"$I${x}", f"$L${x}") for x in pr["rows"]]
rm = metodo(ws, r + 4, "mar", d, it)
ri = metodo(ws, rm[-1] + 3, "int", d, it, cfi)
rd = metodo(ws, ri[-1] + 3, "ide", d, it, cfd)
for i, g in enumerate([(14.00, 15.85, 17.38), (40.89, 34.46, 36.00), (20.22, 20.15, 21.69)]):
    for rows_, p in zip((rm, ri, rd), g):
        exp(ws, f"O{rows_[i]}", p)
r = rd[-1] + 3
title(ws, f"H{r}:M{r}", "c) COMPARAÇÃO COM O MERCADO", bg=CYAN)
header(ws, r + 1, 8, ["PRODUTO", "MARGINAL", "INTEGRAL", "IDEAL", "MERCADO", "MERCADO ÷ IDEAL"])
for i, pr_row in enumerate(pr["rows"]):
    rr = r + 2 + i
    put(ws, f"H{rr}", f"=H{rm[i]}", font="Calibri", size=12)
    put(ws, f"I{rr}", f"=O{rm[i]}", RS, font="Calibri", size=12)
    put(ws, f"J{rr}", f"=O{ri[i]}", RS, font="Calibri", size=12)
    put(ws, f"K{rr}", f"=O{rd[i]}", RS, font="Calibri", size=12)
    put(ws, f"L{rr}", f"=M{pr_row}", RS, font="Calibri", size=12)
    put(ws, f"M{rr}", f"=L{rr}/K{rr}-1", PCT1, font="Calibri", size=12)
box(ws, f"H{r + 2}:M{r + 4}")
r += 7
title(ws, f"H{r}:M{r}", "d) QUEDA NAS VENDAS · NOVO PREÇO PELA ABSORÇÃO IDEAL", bg=CYAN)
header(ws, r + 1, 8, ["PRODUTO", "NOVAS VENDAS", "PREÇO IDEAL", "MERCADO", "SITUAÇÃO"])
novas = [3000, 1000, 2000]
tot2 = r + 5
cfd2 = f"$J${tot2 + 1}"
for i, (q, pr_row) in enumerate(zip(novas, pr["rows"])):
    rr = r + 2 + i
    put(ws, f"H{rr}", f"=H{pr_row}", font="Calibri", size=12)
    put(ws, f"I{rr}", q, INT, font="Calibri", size=12)
    put(ws, f"J{rr}", f"=(I{pr_row}+{cfd2})/((100-({d['pCVV']}+{d['pCVP']}+L{pr_row}))/100)", RS, bold=True, font="Calibri", size=12)
    put(ws, f"K{rr}", f"=M{pr_row}", RS, font="Calibri", size=12)
    put(ws, f"L{rr}", f'=IF(J{rr}<=K{rr},"abaixo do mercado","acima do mercado")', font="Calibri", size=12)
put(ws, f"H{tot2}", "TOTAL", bold=True, font="Calibri", size=12)
put(ws, f"I{tot2}", f"=SUM(I{r + 2}:I{r + 4})", INT, bold=True, font="Calibri", size=12)
box(ws, f"H{r + 2}:L{tot2}")
ws.merge_cells(f"H{tot2 + 1}:I{tot2 + 1}")
put(ws, f"H{tot2 + 1}", "Novo CF unit. ideal = CF ÷ novas vendas", bold=True, font="Calibri", size=11)
put(ws, f"J{tot2 + 1}", f"=ROUND(C{d['CF']}/I{tot2},2)", RS, bold=True, font="Calibri", size=12)
box(ws, f"H{tot2 + 1}:J{tot2 + 1}")
exp(ws, f"J{tot2 + 1}", 8.33)
for i, p in enumerate((22.51, 41.12, 26.82)):
    exp(ws, f"J{r + 2 + i}", p)
resposta(ws, tot2 + 3, "RESPOSTA e)", "Não repassar o aumento: o CF existe com ou sem a venda. Subir o preço reduz o volume, eleva o CF unitário e alimenta a "
         "espiral da morte. Enquanto o preço de mercado deixar margem de contribuição, vale manter o preço e buscar volume.", 50)

for ws in wb:
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
wb.save(OUT)
import json
json.dump(EXPECT, open(OUT + ".expect.json", "w"))
print("ok", len(EXPECT), "checks")
