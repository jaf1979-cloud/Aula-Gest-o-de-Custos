"""Grava os valores calculados (cache) ao lado das fórmulas, para a planilha abrir já com números em qualquer visualizador.
Uso: python3 injeta_cache.py sem_cache.xlsx recalculada_pelo_libreoffice.xlsx saida.xlsx"""
import re, sys, zipfile
from lxml import etree
from openpyxl import load_workbook

SRC, CALC, OUT = sys.argv[1:4]
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
vals = load_workbook(CALC, data_only=True)
zin = zipfile.ZipFile(SRC)
files = {n: zin.read(n) for n in zin.namelist()}
wbx = etree.fromstring(files["xl/workbook.xml"])
rels = etree.fromstring(files["xl/_rels/workbook.xml.rels"])
target = {r.get("Id"): r.get("Target") for r in rels}
n = 0
for s in wbx.find(f"{{{NS}}}sheets"):
    rid = s.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
    path = "xl/" + target[rid]
    ws = vals[s.get("name")]
    root = etree.fromstring(files[path])
    for c in root.iter(f"{{{NS}}}c"):
        f = c.find(f"{{{NS}}}f")
        if f is None:
            continue
        v = ws[c.get("r")].value
        if isinstance(v, (int, float)):
            ve = c.find(f"{{{NS}}}v")
            if ve is None:
                ve = etree.SubElement(c, f"{{{NS}}}v")
            ve.text = repr(float(v)); n += 1
    files[path] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
with zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as z:
    for name in ["[Content_Types].xml"] + [k for k in files if k != "[Content_Types].xml"]:
        z.writestr(name, files[name])
print("valores gravados:", n)
