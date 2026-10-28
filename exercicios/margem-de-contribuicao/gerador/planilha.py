import sys, copy
sys.path.insert(0, __file__.rsplit('/', 1)[0])
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.comments import Comment
from dados import EXS

MON = '#,##0.00;[Red]-#,##0.00'
PCT = '0.00%'
wb = load_workbook(__file__.replace('planilha.py', 'planilha_modelo.xlsx'))
tpl = wb.active
tpl.title = 'Modelo'

def cpstyle(src, dst):
    dst._style = copy.copy(src._style)

for e in EXS:
    ws = wb.copy_worksheet(tpl)
    ws.title = f"Ex{e['n']} {e['ind'].split(' ')[0].split('/')[0]}"
    g = e['gastos']
    R = [v for _, v, c in g if c == 'R'][0]
    cvp_items = [(d, v) for d, v, c in g if c == 'MP'] + [(d, v) for d, v, c in g if c == 'CVP']
    cvv_items = [(d, v) for d, v, c in g if c == 'CVV']
    cf_items = [(d, v) for d, v, c in g if c == 'CF']
    assert len(cvp_items) <= 7 and len(cvv_items) <= 7
    extra = max(0, len(cf_items) - 13)
    if extra:
        ws.insert_rows(42, extra)
        for r in range(42, 42 + extra):
            for col in 'BCD':
                cpstyle(ws[f'{col}41'], ws[f'{col}{r}'])
    lo = 42 + extra
    cf_last = 41 + extra

    ws['B2'] = f"Exercício {e['n']} · {e['empresa']} · {e['cidade']} · mês de {e['mes']}"
    ws['B2'].font = Font(name='Arial', size=10, bold=True)
    ws['B2'].alignment = Alignment(horizontal='center')
    ws['C5'] = R
    for r0, items in ((8, cvp_items), (16, cvv_items)):
        for i, (d, v) in enumerate(items):
            ws[f'B{r0+i}'] = d; ws[f'C{r0+i}'] = v
    for i, (d, v) in enumerate(cf_items):
        ws[f'B{29+i}'] = d; ws[f'C{29+i}'] = v
    ws['C7'] = '=SUM(C8:C14)'
    ws['C15'] = '=SUM(C16:C22)'
    ws['C24'] = '=C7+C15'
    ws['C26'] = '=C5-C24'
    ws['C28'] = f'=SUM(C29:C{cf_last})'
    ws[f'C{lo}'] = f'=C26-C28'
    rows_d = [5, 7, 15, 24, 26, 28, lo] + list(range(8, 8 + len(cvp_items))) + list(range(16, 16 + len(cvv_items))) + list(range(29, 29 + len(cf_items)))
    for r in range(5, lo + 1):
        if ws[f'C{r}'].value is not None:
            ws[f'C{r}'].number_format = MON
    for r in rows_d:
        ws[f'D{r}'] = f'=C{r}/$C$5'
        ws[f'D{r}'].number_format = PCT
    ws['C8'].comment = Comment('Matéria-prima: entra em R$ pela ficha técnica (coluna MP); por isso sai do CVP % da tabela.', 'Prof.')

    # tabela de produtos
    ws['F6'] = f"CUSTO PELA MARGEM DE CONTRIBUIÇÃO · {e['empresa']}"
    for i, p in enumerate(e['prods']):
        r = 10 + i
        ws[f'F{r}'] = p['nome']; ws[f'G{r}'] = p['q']; ws[f'H{r}'] = p['pv']; ws[f'I{r}'] = p['mp']
        ws[f'J{r}'] = '=($C$7-$C$8)/$C$5'
        ws[f'K{r}'] = f'=H{r}*J{r}'
        ws[f'L{r}'] = '=$C$15/$C$5'
        ws[f'M{r}'] = f'=H{r}*L{r}'
        ws[f'N{r}'] = f'=I{r}+K{r}+M{r}'
        ws[f'O{r}'] = f'=H{r}-N{r}'
        ws[f'P{r}'] = f'=O{r}/H{r}'
        ws[f'Q{r}'] = f'=O{r}*G{r}'
        ws[f'R{r}'] = f'=Q{r}/$Q$17'
        ws[f'S{r}'] = f'=$C$28*R{r}/G{r}'
        ws[f'T{r}'] = f'=(I{r}+S{r})/(1-(J{r}+L{r}))'
        ws[f'U{r}'] = p['L'] / 100
        ws[f'V{r}'] = f'=(I{r}+S{r})/(1-(J{r}+L{r}+U{r}))'
        ws[f'G{r}'].number_format = '#,##0'
        for c in 'HIKMNOQSTV':
            ws[f'{c}{r}'].number_format = MON
        for c in 'JLPRU':
            ws[f'{c}{r}'].number_format = PCT
    ws['Q17'] = '=SUM(Q10:Q16)'; ws['Q17'].number_format = MON
    ws['R17'] = '=SUM(R10:R16)'; ws['R17'].number_format = PCT

    # legenda e conferências
    ws['F19'] = 'Produtos:'; ws['F19'].font = Font(name='Arial', size=10, bold=True)
    for i, p in enumerate(e['prods']):
        ws[f'F{20+i}'] = f"{p['nome']} = {p['desc']}"; ws[f'F{20+i}'].font = Font(name='Arial', size=10)
    r = 21 + len(e['prods'])
    ws[f'F{r}'] = 'Conferência: Σ MC total dos produtos − MC do DRE (deve ser 0)'; ws[f'F{r}'].font = Font(name='Arial', size=10, italic=True)
    ws[f'Q{r}'] = '=Q17-C26'; ws[f'Q{r}'].number_format = MON
    ws[f'F{r+1}'] = 'Conferência: Σ quantidade × preço − receita (deve ser 0)'; ws[f'F{r+1}'].font = Font(name='Arial', size=10, italic=True)
    ws[f'Q{r+1}'] = '=SUMPRODUCT(G10:G16,H10:H16)-C5'; ws[f'Q{r+1}'].number_format = MON
    ws[f'F{r+3}'] = ('Fórmulas: CVP% = (CVP total − MP) ÷ Receita · CF unit = CF total × MC total % ÷ quantidade · '
                     'Custo = (MP + CF unit) ÷ [1 − (CVP% + CVV%)] · PV = (MP + CF unit) ÷ [1 − (CVP% + CVV% + Lucro%)]')
    ws[f'F{r+3}'].font = Font(name='Arial', size=9, italic=True)
    for c in 'GIJKLMNOPRSTUV':
        ws.column_dimensions[c].width = 13
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['Q'].width = 16
    ws.column_dimensions['T'].width = 15
    ws.column_dimensions['V'].width = 14

wb.remove(tpl)
wb.save(sys.argv[1])
print([s.title for s in wb])
