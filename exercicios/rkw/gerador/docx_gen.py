"""Gera os arquivos Word: enunciados (para os alunos) e resolução passo a passo.
Uso: python3 docx_gen.py ../rkw-exercicios-enunciados.docx ../rkw-exercicios-resolucao.docx"""
import re
import sys
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor
from gen import build

NAVY = RGBColor(0x0B, 0x2A, 0x63)
TEAL = RGBColor(0x0B, 0x7D, 0x72)
GREY = RGBColor(0x4A, 0x58, 0x66)


def shade(el, fill):
    pr = el.get_or_add_tcPr() if hasattr(el, "get_or_add_tcPr") else el._p.get_or_add_pPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"), s.set(qn("w:color"), "auto"), s.set(qn("w:fill"), fill)
    pr.append(s)


def border_left(par, color):
    pPr = par._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for k, v in (("val", "single"), ("sz", "18"), ("space", "6"), ("color", color)):
        left.set(qn("w:" + k), v)
    b.append(left)
    pPr.append(b)


def runs(par, text, size=10.5, color=None, italic=False):
    for i, part in enumerate(re.split(r"</?b>", text)):
        if not part:
            continue
        r = par.add_run(part)
        r.bold = i % 2 == 1
        r.font.size = Pt(size)
        r.italic = italic
        if color:
            r.font.color.rgb = color


def para(doc, text, **kw):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    runs(p, text, **kw)
    return p


def table(doc, title, head, rows, right, bold_rows):
    if title:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title.upper())
        p.paragraph_format.keep_with_next = True
        r.bold, r.font.size, r.font.color.rgb = True, Pt(9.5), NAVY
    ncol = len(head)
    fs = 9 if ncol <= 5 else (8 if ncol <= 7 else 7)
    t = doc.add_table(rows=1, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]
        c.text = ""
        runs(c.paragraphs[0], f"<b>{h}</b>", size=fs, color=NAVY)
        shade(c._tc, "EAF1F5")
        if i in right:
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for k, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            txt = f"<b>{v}</b>" if (k in bold_rows or i == 0) and v else v
            runs(cells[i].paragraphs[0], txt, size=fs)
            if i in right:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if k in bold_rows:
                shade(cells[i]._tc, "FFF3C4")
    total = 17.8
    first = 3.2 if ncol > 4 else 4.5
    ws = [first] + [(total - first) / (ncol - 1)] * (ncol - 1)
    t.autofit = False
    for row in t.rows:
        for i, c in enumerate(row.cells):
            c.width = Cm(ws[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render(doc, blocks):
    for b in blocks:
        k = b[0]
        if k == "h":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(b[1])
            r.bold, r.font.size, r.font.color.rgb = True, Pt(12), NAVY
        elif k == "p":
            para(doc, b[1])
        elif k == "eq":
            p = para(doc, b[1], size=10)
            p.paragraph_format.left_indent = Cm(0.6)
        elif k == "list":
            for it in b[1]:
                p = para(doc, "➔  " + it, size=10)
                p.paragraph_format.left_indent = Cm(0.6)
        elif k == "table":
            table(doc, *b[1:])
        elif k == "note":
            p = para(doc, b[1], size=9.5, color=GREY, italic=True)
        elif k == "interp":
            p = para(doc, b[1], size=10)
            p.paragraph_format.left_indent = Cm(0.3)
            border_left(p, "0F9E93")
            shade(p, "E2F4F2")


def novo(titulo, sub):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.font.size = Pt(10.5)
    s = doc.sections[0]
    s.page_height, s.page_width = Cm(29.7), Cm(21)
    s.orientation = WD_ORIENT.PORTRAIT
    for side in ("left_margin", "right_margin"):
        setattr(s, side, Cm(1.6))
    s.top_margin = s.bottom_margin = Cm(1.6)
    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("UTFPR · Gestão de Custos · Prof. José Ângelo Ferreira")
    r.font.size, r.font.color.rgb = Pt(8), GREY
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("UTFPR · GESTÃO DE CUSTOS")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(10), GREY
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(titulo)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(15), NAVY
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(sub)
    r.font.size, r.font.color.rgb = Pt(10), GREY
    return doc


def quebra(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


REGRA = ("<b>Respostas:</b> não utilizar ponto para milhar e usar vírgula para as decimais nos valores monetários "
         "(ex.: 100200,00). Utilizar 2 casas após a vírgula para valores monetários e numéricos (ex.: 200,00; 0,19).")

EX = build()

# Enunciados
doc = novo("Custeio RKW – Lista de Exercícios", "Matriz RKW, custo por minuto, custo RKW, custo total e preço de venda")
para(doc, "Nome: ____________________________________   RA: ___________   Data: ___/___/_____", size=10)
para(doc, REGRA, size=10)
for k, e in enumerate(EX):
    if k:
        quebra(doc)
    d = e["d"]
    p = doc.add_paragraph()
    r = p.add_run(f"Exercício {d['n']} · {d['setor']} – {d['empresa']}")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(13), NAVY
    render(doc, e["enun"])
doc.save(sys.argv[1])

# Resolução
doc = novo("Custeio RKW – Resolução Passo a Passo", "Exercícios 1 a 5 · uso do professor")
para(doc, "Método (conforme os slides): 1º rateio pelos direcionadores → 2º rateio dos centros de apoio pelo nº de "
          "funcionários → custo por minuto = custo do centro ÷ (horas × 60 × funcionários) → custo RKW = Σ tempo × "
          "custo/min → custo total = (MP + RKW) ÷ [(100 − (CVV + CVP)) ÷ 100] → preço = (MP + RKW) ÷ "
          "[(100 − (CVV + CVP + L)) ÷ 100].", size=10)
for k, e in enumerate(EX):
    if k:
        quebra(doc)
    d = e["d"]
    p = doc.add_paragraph()
    r = p.add_run(f"Exercício {d['n']} · {d['setor']} – {d['empresa']}")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(13), NAVY
    render(doc, e["sol"])
doc.save(sys.argv[2])
print("ok")
