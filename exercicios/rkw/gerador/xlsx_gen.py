"""Preenche o modelo "Matriz RKW.xlsx" do professor: uma aba por exercício, com fórmulas do Excel.
Uso: python3 xlsx_gen.py modelo.xlsx ../rkw-exercicios-resolucao.xlsx

Os dados de entrada (custos, direcionadores, horas, tempos, MP, CVV, CVP, lucro) ficam em tabelas próprias;
a Matriz RKW, o custo por minuto e o custo dos produtos são fórmulas que apontam para essas tabelas.
O estilo de cada célula é copiado da célula equivalente do modelo."""
import sys
from copy import copy
import openpyxl
from openpyxl.utils import get_column_letter as L
from dados import EXS

MOD, OUT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(MOD)
T = wb["RKW"]

RS = '_-"R$"* #,##0.00_-;\\-"R$"* #,##0.00_-;_-"R$"* "-"??_-;_-@'
N0, N2 = "#,##0", "#,##0.00"
DRV_ROW = {"Funcionários": "Nº FUNCIONÁRIOS", "Área": "ÁREA OCUPADA (M²)", "Potência": "POTÊNCIA INSTALADA (KW)",
           "Requisições": "Nº REQUISIÇÕES", "Ordens": "Nº ORDENS DE MANUTENÇÃO"}
ITEM_DRV_TXT = {"Direto": "Direto", "Área": "M²", "Potência": "Potência Inst.", "Requisições": "Nº Requisições",
                "Ordens": "Nº Ordens", "Funcionários": "Nº Func"}


def put(ws, ref, value, src, fmt=None):
    c = ws[ref]
    c.value = value
    s = T[src]
    c._style = copy(s._style)
    if fmt:
        c.number_format = fmt
    return c


def bar(ws, row, c1, c2, text, src):
    for c in range(c1, c2 + 1):
        put(ws, f"{L(c)}{row}", text if c == c1 else None, src)
    if c2 > c1:
        ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)


def sheet(d):
    ws = wb.create_sheet(f"Ex{d['n']} {d['empresa'].split()[0]}")
    ws.sheet_view.showGridLines = False
    C, apoio = d["centros"], d["apoio"]
    prod = [c for c in C if c not in apoio]
    n = len(C)
    col = {c: 3 + i for i, c in enumerate(C)}       # C, D, E...
    TOT = 3 + n                                       # coluna do total
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 46
    for i in range(n + 1):
        ws.column_dimensions[L(3 + i)].width = 28 if i == 0 else 17
    ws["B1"] = f"Exercício {d['n']} – {d['empresa']}"
    ws["B1"].font = copy(T["B3"].font)

    # ---------- tabelas de dados (entradas) ----------
    # custos do mês: à direita da matriz
    kc = TOT + 2
    for c in range(kc, kc + 3):
        ws.column_dimensions[L(c)].width = 22 if c == kc else 16
    bar(ws, 2, kc, kc + 2, "CUSTOS DO MÊS", "C26")
    for j, h in enumerate(("CUSTO", "VALOR", "DIRECIONADOR")):
        put(ws, f"{L(kc + j)}3", h, "B27")
    item_ref = {}
    for i, (nome, valor, drv) in enumerate(d["itens"]):
        r = 4 + i
        put(ws, f"{L(kc)}{r}", nome, "B28")
        put(ws, f"{L(kc + 1)}{r}", valor, "C28", N2)
        put(ws, f"{L(kc + 2)}{r}", ITEM_DRV_TXT[drv], "C28")
        item_ref[nome] = f"${L(kc + 1)}${r}"
    r = 4 + len(d["itens"])
    put(ws, f"{L(kc)}{r}", "TOTAL", "B27")
    put(ws, f"{L(kc + 1)}{r}", f"=SUM({L(kc + 1)}4:{L(kc + 1)}{r - 1})", "B27", N2)
    put(ws, f"{L(kc + 2)}{r}", None, "B27")

    # ---------- MATRIZ RKW ----------
    bar(ws, 2, 3, TOT - 1, "SEÇÕES", "C2")
    put(ws, "B3", "CUSTOS/DIRECIONADORES", "B3")
    for c in C:
        put(ws, f"{L(col[c])}3", c.upper(), "C3")
    put(ws, f"{L(TOT)}3", "TOTAL", "C3")

    # linhas da matriz: salários (direcionador nº func.) + um par de linhas por item rateado
    drv_tab_row = {}   # preenchido abaixo; precisamos das linhas da tabela de direcionadores
    salarios = [(nm, v) for nm, v, k in d["itens"] if k == "Direto"]
    outros = [(nm, v, k) for nm, v, k in d["itens"] if k != "Direto"]
    pares = 1 + len(outros)
    row1 = 4 + 2 * pares            # 1º rateio
    sec_rows = list(range(row1 + 1, row1 + 1 + len(apoio)))
    rtot = sec_rows[-1] + 1
    # tabela de direcionadores (abaixo da matriz)
    dt = rtot + 4
    drvs = [k for k in ("Funcionários", "Área", "Potência", "Requisições", "Ordens") if k in d["drivers"]]
    put(ws, f"B{dt}", "DIRECIONADORES DE RATEIO", "B25")
    bar(ws, dt + 1, 3, TOT - 1, "SEÇÕES", "C26")
    put(ws, f"B{dt + 2}", "DIRECIONADORES", "B27")
    for c in C:
        put(ws, f"{L(col[c])}{dt + 2}", c.upper(), "C27")
    for i, k in enumerate(drvs + ["Horas"]):
        r = dt + 3 + i
        put(ws, f"B{r}", DRV_ROW.get(k, "HORAS TRABALHADAS/MÊS (POR FUNC.)"), "B28")
        for c in C:
            v = d["drivers"][k][c] if k != "Horas" else (None if c in apoio else d["horas"])
            put(ws, f"{L(col[c])}{r}", v, "C28", N0)
        drv_tab_row[k] = r
    dt_end = dt + 3 + len(drvs)

    def dref(k, c, abs_=False):
        return f"{L(col[c])}{drv_tab_row[k]}" if not abs_ else f"${L(col[c])}${drv_tab_row[k]}"

    lab = [("B4", "C4"), ("B6", "C6")]
    # par 0: salários
    put(ws, "B4", DRV_ROW["Funcionários"], lab[0][0])
    for c in C:
        put(ws, f"{L(col[c])}4", f"={dref('Funcionários', c)}", lab[0][1], N0)
    put(ws, f"{L(TOT)}4", None, lab[0][1])
    put(ws, "B5", "SALÁRIOS (DIRETO)", "B5")
    for c in C:
        nome = next(nm for nm, _ in salarios if d["salario_de"][nm] == c)
        put(ws, f"{L(col[c])}5", f"={item_ref[nome]}", "C5", RS)
    put(ws, f"{L(TOT)}5", f"=SUM(C5:{L(TOT - 1)}5)", "C5", RS)
    for i, (nome, v, k) in enumerate(outros):
        rd, rc = 6 + 2 * i, 7 + 2 * i
        bl, cl = lab[(i + 1) % 2]
        put(ws, f"B{rd}", DRV_ROW[k], bl)
        for c in C:
            put(ws, f"{L(col[c])}{rd}", f"={dref(k, c)}", cl, N0)
        put(ws, f"{L(TOT)}{rd}", f"=SUM(C{rd}:{L(TOT - 1)}{rd})", cl, N0)
        put(ws, f"B{rc}", nome.upper(), "B5")
        for c in C:
            put(ws, f"{L(col[c])}{rc}", f"={item_ref[nome]}*{L(col[c])}{rd}/${L(TOT)}${rd}", "C5", RS)
        put(ws, f"{L(TOT)}{rc}", f"=SUM(C{rc}:{L(TOT - 1)}{rc})", "C5", RS)
    custo_rows = [5] + [7 + 2 * i for i in range(len(outros))]
    put(ws, f"B{row1}", "1º RATEIO  ∑", "B18")
    for c in C + [None]:
        cc = L(col[c]) if c else L(TOT)
        put(ws, f"{cc}{row1}", "=" + "+".join(f"{cc}{r}" for r in custo_rows), "C18", RS)
    # 2º rateio em degraus, pelo nº de funcionários de quem recebe
    for i, a in enumerate(apoio):
        r = sec_rows[i]
        recebem = [c for c in C if c not in apoio[: i + 1]]
        put(ws, f"B{r}", f"2º RATEIO – {a.upper()}  ∑", "B19")
        # valor a distribuir = 1º rateio + o que recebeu dos apoios anteriores
        a_val = f"({L(col[a])}{row1}" + "".join(f"+{L(col[a])}{sec_rows[j]}" for j in range(i)) + ")"
        base = "(" + "+".join(dref("Funcionários", c, True) for c in recebem) + ")"
        for c in C:
            if c == a:
                f = f"=-{a_val}"
            elif c in recebem:
                f = f"={a_val.replace(L(col[a]), '$' + L(col[a]))}*{dref('Funcionários', c, True)}/{base}"
            else:
                f = None
            put(ws, f"{L(col[c])}{r}", f, "C19", RS)
        put(ws, f"{L(TOT)}{r}", f"=SUM(C{r}:{L(TOT - 1)}{r})", "C19", RS)
    put(ws, f"B{rtot}", None, "B20")
    put(ws, f"{L(col[apoio[-1]])}{rtot}", "∑ TOTAL", "C20")
    for c in prod + [None]:
        cc = L(col[c]) if c else L(TOT)
        put(ws, f"{cc}{rtot}", f"={cc}{row1}+" + "+".join(f"{cc}{r}" for r in sec_rows), "D20", RS)
    tot_ref = {c: f"{L(col[c])}{rtot}" for c in prod}

    # ---------- tempo de processo e MP ----------
    r0 = dt_end + 3
    bar(ws, r0, 3, 3 + len(prod), "TEMPO DE PROCESSO DOS PRODUTOS (MIN/UN)", "C36")
    put(ws, f"C{r0 + 1}", "PRODUTO  / SEÇÃO", "C37")
    for j, c in enumerate(prod):
        put(ws, f"{L(4 + j)}{r0 + 1}", c.upper(), "D37")
    t_ref = {}
    for i, p in enumerate(d["produtos"]):
        r = r0 + 2 + i
        put(ws, f"C{r}", p["nome"].upper(), "C38")
        for j, c in enumerate(prod):
            put(ws, f"{L(4 + j)}{r}", p["t"][c], "D38", N0)
            t_ref[(i, c)] = f"{L(4 + j)}{r}"
    r0 = r0 + 2 + len(d["produtos"]) + 2
    bar(ws, r0, 3, 7, "MATÉRIA-PRIMA / CUSTOS VARIÁVEIS (%) / MARGEM DE LUCRO (%)", "C42")
    for j, h in enumerate(("PRODUTO", "MAT PRIM (R$)", "CVV (%)", "CVP (%)", "LUCRO (%)")):
        put(ws, f"{L(3 + j)}{r0 + 1}", h, "C43" if j == 0 else "D43")
    mp_ref = {}
    for i, p in enumerate(d["produtos"]):
        r = r0 + 2 + i
        put(ws, f"C{r}", p["nome"].upper(), "C44")
        for j, v in enumerate((p["mp"], p["cvv"], p["cvp"], d["lucro"])):
            put(ws, f"{L(4 + j)}{r}", v, "D44", N2)
        mp_ref[i] = r
    # ---------- custo minuto ----------
    r0 = r0 + 2 + len(d["produtos"]) + 2
    put(ws, f"C{r0}", "CUSTO MINUTO DAS SEÇÕES", "D47")
    ws.merge_cells(f"C{r0}:D{r0}")
    put(ws, f"E{r0}", "custo/min arredondado em 2 casas (regra das respostas)", "D50")
    ws[f"E{r0}"].font = copy(T["E50"].font)
    ws[f"E{r0}"].font = openpyxl.styles.Font(name=T["E50"].font.name, sz=10, i=True, color="FF5A6878")
    r = r0 + 2
    min_ref = {}
    for c in prod:
        put(ws, f"C{r}", c.upper(), "D49")
        put(ws, f"D{r}", None, "D49")
        ws.merge_cells(f"C{r}:D{r}")
        rows = [("TEMPO TRABALHADO (MIN)", f"={dref('Horas', c)}*60", N0),
                ("NÚMERO DE FUNCIONÁRIOS", f"={dref('Funcionários', c)}", N0),
                ("CUSTO TOTAL DA SEÇÃO", f"={tot_ref[c]}", RS),
                ("CUSTO MINUTO", f"=ROUND(D{r + 3}/(D{r + 1}*D{r + 2}),2)", N2)]
        for k, (lb, f, fmt) in enumerate(rows):
            put(ws, f"C{r + 1 + k}", lb, "D50")
            put(ws, f"D{r + 1 + k}", f, "E50", fmt)
        min_ref[c] = f"$D${r + 4}"
        r += 6
    # ---------- custo dos produtos ----------
    r0 = r + 1
    bar(ws, r0, 3, 6, "CUSTO RKW / CUSTO TOTAL / PREÇO DE VENDA", "C68")
    r = r0 + 2
    for i, p in enumerate(d["produtos"]):
        bar(ws, r, 3, 6, p["nome"].upper(), "C70")
        for j, h in enumerate(("CUSTO SEÇÕES", "TEMPO PROCESSO", "CUSTO MINUTO", "CUSTO RKW")):
            put(ws, f"{L(3 + j)}{r + 1}", h, "C71" if j == 0 else "D71")
        for k, c in enumerate(prod):
            rr = r + 2 + k
            put(ws, f"C{rr}", f"CUSTO {c.upper()}", "C72")
            put(ws, f"D{rr}", f"={t_ref[(i, c)]}", "D72", N0)
            put(ws, f"E{rr}", f"={min_ref[c]}", "D72", N2)
            put(ws, f"F{rr}", f"=D{rr}*E{rr}", "D72", RS)
        rr = r + 2 + len(prod)
        m = mp_ref[i]
        put(ws, f"C{rr}", "CUSTO TOTAL RKW", "C75")
        put(ws, f"D{rr}", f"=SUM(F{r + 2}:F{rr - 1})", "D75", RS)
        put(ws, f"C{rr + 1}", "CUSTO TOTAL", "C75")
        put(ws, f"D{rr + 1}", f"=(D{m}+D{rr})/((100-(E{m}+F{m}))/100)", "D75", RS)
        put(ws, f"C{rr + 2}", "PREÇO DE VENDAS", "C75")
        put(ws, f"D{rr + 2}", f"=(D{m}+D{rr})/((100-(E{m}+F{m}+G{m}))/100)", "D75", RS)
        r = rr + 4
    ws.print_area = f"A1:{L(kc + 2)}{r}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


for d in EXS:
    sheet(d)
wb.remove(T)
wb.save(OUT)
print("ok")
